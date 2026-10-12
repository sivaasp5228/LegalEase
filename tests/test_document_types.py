"""Integration tests verifying all 5 supported document types across all exporters."""

import io
import pytest
from docx import Document
from document_generator.txt_generator import generate_txt
from document_generator.docx_generator import generate_docx
from document_generator.pdf_generator import generate_pdf
from document_generator.html_preview import generate_html_preview


DOC_TYPE_SAMPLES = [
    (
        "Agreement",
        (
            "COMMERCIAL PARTNERSHIP AGREEMENT\n\n"
            "This Partnership Agreement is made effective October 1, 2026.\n\n"
            "PARTIES: Alpha Corp and Beta Enterprises\n\n"
            "| Term | Specification |\n"
            "| Profit Distribution | 50% Mutual Distribution |\n"
            "| Initial Contribution | $50,000 USD |\n\n"
            "SECTION 1: PURPOSE\n"
            "Collaborative software engineering services.\n\n"
            "BY:\n"
            "NAME: Alice Stone\n"
            "TITLE: Managing Partner"
        ),
    ),
    (
        "Contract",
        (
            "INDEPENDENT CONTRACTOR CONTRACT\n\n"
            "Effective Date: October 1, 2026\n\n"
            "PARTIES: Tech Services Corp (Client) and DevCraft LLC (Contractor)\n\n"
            "| Deliverable | Milestone Payment |\n"
            "| Backend Architecture | $15,000 |\n"
            "| Frontend Delivery | $15,000 |\n\n"
            "SECTION 1: SCOPE OF WORK\n"
            "Contractor agrees to deliver cloud services.\n\n"
            "BY:\n"
            "NAME: Bob Vance\n"
            "TITLE: President"
        ),
    ),
    (
        "NDA",
        (
            "MUTUAL NON-DISCLOSURE AGREEMENT\n\n"
            "Effective Date: October 1, 2026\n\n"
            "PARTIES: Disclosing Party and Receiving Party\n\n"
            "| Term | Duration |\n"
            "| Confidentiality Period | 3 Years |\n\n"
            "SECTION 1: TRADE SECRETS\n"
            "Recipient must safeguard confidential source code.\n\n"
            "BY:\n"
            "NAME: Carol Danvers\n"
            "TITLE: Chief Legal Officer"
        ),
    ),
    (
        "Lease Agreement",
        (
            "RESIDENTIAL LEASE AGREEMENT\n\n"
            "Effective Date: October 1, 2026\n\n"
            "PARTIES: Landlord LLC and Tenant John Smith\n\n"
            "| Provision | Agreed Term |\n"
            "| Monthly Rent | $2,200 USD |\n"
            "| Security Deposit | $2,200 USD |\n\n"
            "SECTION 1: DEMISED PREMISES\n"
            "Apartment 4B, 742 Evergreen Terrace.\n\n"
            "BY:\n"
            "NAME: David Miller\n"
            "TITLE: Property Manager"
        ),
    ),
    (
        "Employment Offer Letter",
        (
            "EMPLOYMENT OFFER LETTER\n\n"
            "Effective Date: October 1, 2026\n\n"
            "PARTIES: NextGen Robotics (Employer) and Jane Doe (Candidate)\n\n"
            "| Compensation Item | Details |\n"
            "| Base Annual Salary | $165,000 USD |\n"
            "| Incentive Bonus | 15% Annual Target |\n\n"
            "SECTION 1: TITLE AND RESPONSIBILITIES\n"
            "Candidate is offered the role of Lead AI Engineer.\n\n"
            "BY:\n"
            "NAME: Elena Rostova\n"
            "TITLE: VP of People"
        ),
    ),
]


@pytest.mark.parametrize("doc_type, sample_text", DOC_TYPE_SAMPLES)
def test_all_document_types_txt_export(doc_type, sample_text):
    """Verify TXT exporter handles all 5 document types."""
    data = generate_txt(sample_text)
    assert isinstance(data, bytes)
    assert len(data) > 0
    decoded = data.decode("utf-8").upper()
    keywords = [doc_type.upper(), "NON-DISCLOSURE"] if doc_type == "NDA" else [doc_type.upper()]
    assert any(k in decoded for k in keywords)


@pytest.mark.parametrize("doc_type, sample_text", DOC_TYPE_SAMPLES)
def test_all_document_types_docx_export(doc_type, sample_text):
    """Verify DOCX exporter produces valid Word documents with tables for all 5 document types."""
    docx_bytes = generate_docx(sample_text, title=doc_type.upper())
    assert isinstance(docx_bytes, bytes)
    assert len(docx_bytes) > 5000
    doc = Document(io.BytesIO(docx_bytes))
    assert len(doc.paragraphs) > 0
    assert len(doc.tables) >= 1


@pytest.mark.parametrize("doc_type, sample_text", DOC_TYPE_SAMPLES)
def test_all_document_types_pdf_export(doc_type, sample_text):
    """Verify PDF exporter produces valid PDF files starting with %PDF for all 5 document types."""
    pdf_bytes = generate_pdf(sample_text, title=doc_type.upper())
    assert isinstance(pdf_bytes, bytes)
    assert pdf_bytes.startswith(b"%PDF")
    assert len(pdf_bytes) > 3000


@pytest.mark.parametrize("doc_type, sample_text", DOC_TYPE_SAMPLES)
def test_all_document_types_html_preview(doc_type, sample_text):
    """Verify HTML preview produces valid styled parchment containers for all 5 document types."""
    html_preview = generate_html_preview(sample_text)
    assert '<div class="legal-parchment">' in html_preview
    assert '<table' in html_preview
