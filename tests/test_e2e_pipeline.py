"""Comprehensive end-to-end integration test suite for LegalEase."""

import io
from unittest.mock import patch
import pytest
from fastapi.testclient import TestClient
from docx import Document

from app.backend.main import app
from document_generator.txt_generator import generate_txt
from document_generator.docx_generator import generate_docx
from document_generator.pdf_generator import generate_pdf
from document_generator.html_preview import generate_html_preview

client = TestClient(app)

SAMPLE_DOCUMENTS = [
    (
        "Agreement",
        "Apex Corp (Vendor) and Beta Dynamics (Client)",
        "Net 30 terms, milestone delivery, Delaware governing law.",
        "2026-10-01",
        (
            "COMMERCIAL SERVICES AGREEMENT\n\n"
            "This Agreement is entered into on October 1, 2026.\n\n"
            "PARTIES: Apex Corp and Beta Dynamics\n\n"
            "| Term | Details |\n"
            "| Payment Schedule | Net 30 Days |\n"
            "| Governing Law | Delaware |\n\n"
            "SECTION 1: SERVICES\n"
            "Vendor shall provide enterprise cloud services.\n\n"
            "BY:\n"
            "NAME: Alice Stone\n"
            "TITLE: VP Operations"
        ),
    ),
    (
        "NDA",
        "Disclosing Corp and Receiving LLC",
        "Mutual 2-year non-disclosure, trade secrets, no reverse engineering.",
        "2026-10-01",
        (
            "MUTUAL NON-DISCLOSURE AGREEMENT\n\n"
            "Effective Date: October 1, 2026\n\n"
            "PARTIES: Disclosing Corp and Receiving LLC\n\n"
            "| Item | Specification |\n"
            "| Duration | 2 Years |\n\n"
            "SECTION 1: CONFIDENTIALITY\n"
            "* Receiving party will not disclose proprietary code.\n"
            "* Materials to be returned upon request.\n\n"
            "BY:\n"
            "NAME: Bob Smith\n"
            "TITLE: Chief Legal Officer"
        ),
    ),
]


@pytest.mark.parametrize("doc_type, parties, terms, dates, mock_content", SAMPLE_DOCUMENTS)
def test_full_e2e_lifecycle(doc_type, parties, terms, dates, mock_content):
    """Test full pipeline: API generation -> TXT -> DOCX -> PDF -> HTML preview."""
    payload = {
        "document_type": doc_type,
        "parties": parties,
        "terms": terms,
        "dates": dates,
    }

    with patch("app.backend.routes.GeminiDocumentGenerator.generate_document", return_value=mock_content):
        # 1. API Request
        response = client.post("/generate", json=payload)
        assert response.status_code == 200
        data = response.json()
        assert data["status"] == "success"
        assert data["document_type"] == doc_type
        content = data["content"]

        # 2. TXT Export
        txt_bytes = generate_txt(content)
        assert isinstance(txt_bytes, bytes)
        assert len(txt_bytes) > 0

        # 3. DOCX Export
        docx_bytes = generate_docx(content, title=doc_type.upper())
        assert isinstance(docx_bytes, bytes)
        assert len(docx_bytes) > 5000
        doc = Document(io.BytesIO(docx_bytes))
        assert len(doc.tables) >= 1
        assert len(doc.paragraphs) > 0

        # 4. PDF Export
        pdf_bytes = generate_pdf(content, title=doc_type.upper())
        assert isinstance(pdf_bytes, bytes)
        assert pdf_bytes.startswith(b"%PDF")
        assert len(pdf_bytes) > 2000

        # 5. HTML Preview
        html_output = generate_html_preview(content)
        assert '<div class="legal-parchment">' in html_output
        assert '<table' in html_output


def test_e2e_edge_case_unicode_and_emojis():
    """Verify document exports handle emojis and uncommon unicode without crashing."""
    unicode_doc = (
        "EMPLOYMENT OFFER LETTER \u2696\ufe0f\n\n"
        "Effective Date: October 1, 2026\n\n"
        "PARTIES: Tech Corp \U0001f680 and Jane Doe\n\n"
        "| Benefit | Value |\n"
        "| Salary | $150,000 \u20ac |\n"
        "| Perks | Caf\u00e9, R\u00e9sum\u00e9 assistance |\n\n"
        "SECTION 1: AT-WILL \u00a7\n"
        "\u201cAt-will employment applies at all times.\u201d\n\n"
        "BY:\n"
        "NAME: HR Director"
    )

    # Test all 4 format pipelines
    txt = generate_txt(unicode_doc)
    assert isinstance(txt, bytes)

    docx = generate_docx(unicode_doc, title="OFFER LETTER")
    assert isinstance(docx, bytes)

    pdf = generate_pdf(unicode_doc, title="OFFER LETTER")
    assert isinstance(pdf, bytes)
    assert pdf.startswith(b"%PDF")

    html = generate_html_preview(unicode_doc)
    assert '<div class="legal-parchment">' in html


def test_e2e_validation_errors():
    """Verify detailed validation errors for bad inputs."""
    # Empty body
    r1 = client.post("/generate", json={})
    assert r1.status_code == 422

    # Whitespace only
    r2 = client.post("/generate", json={
        "document_type": "NDA",
        "parties": "   ",
        "terms": "Valid terms here",
        "dates": "2026-10-01",
    })
    assert r2.status_code == 422
    assert "parties" in str(r2.json()).lower()

    # Missing terms
    r3 = client.post("/generate", json={
        "document_type": "NDA",
        "parties": "Valid Party",
        "terms": "   ",
        "dates": "2026-10-01",
    })
    assert r3.status_code == 422
    assert "terms" in str(r3.json()).lower()
